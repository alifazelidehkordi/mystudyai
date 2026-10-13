import json
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as SlideInches

from mystudyai.cli import main
from mystudyai.ingest import import_document


def load(folder):
    return json.loads((folder / "document.json").read_text(encoding="utf-8"))


def make_image(path: Path):
    Image.new("RGB", (20, 20), "white").save(path)


def test_markdown_import_and_idempotence(tmp_path):
    file = tmp_path / "lesson.md"
    file.write_text("# Membrane potential\n\nThe Na+/K+ pump uses ATP.\n", encoding="utf-8")
    output = import_document(file, tmp_path / "library")
    obj = load(output)
    assert [block["location"] for block in obj["blocks"]] == ["line 1", "line 3"]
    assert "Membrane potential" in (output / "notes.md").read_text()
    assert import_document(file, tmp_path / "library") == output


def test_pdf_extracts_text_page_citation_and_image(tmp_path):
    image = tmp_path / "micrograph.png"
    make_image(image)
    pdf_file = tmp_path / "physiology.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((50, 50), "Arteries carry blood away from the heart")
    page.insert_image(fitz.Rect(30, 70, 120, 160), filename=str(image))
    pdf.new_page()  # A page without selectable text should be flagged
    pdf.save(pdf_file)
    pdf.close()
    output = import_document(pdf_file, tmp_path / "library")
    obj = load(output)
    assert any("Arteries carry blood" in block["text"] and block["location"] == "page 1" for block in obj["blocks"])
    assert any(figure["location"] == "page 1" and (output / figure["path"]).is_file() for figure in obj["figures"])
    assert any("page 2" in warning and "OCR" in warning for warning in obj["warnings"])
    assert "Source: physiology.pdf, page 1" in (output / "notes.md").read_text()


def test_powerpoint_text_table_image_and_slides(tmp_path):
    image = tmp_path / "cardiac.png"
    make_image(image)
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    textbox = slide.shapes.add_textbox(SlideInches(1), SlideInches(1), SlideInches(4), SlideInches(1))
    textbox.text = "Systole and diastole"
    table = slide.shapes.add_table(1, 2, SlideInches(1), SlideInches(2), SlideInches(4), SlideInches(1)).table
    table.cell(0, 0).text = "Aorta"
    table.cell(0, 1).text = "Vessel"
    slide.shapes.add_picture(str(image), SlideInches(1), SlideInches(3))
    file = tmp_path / "heart.pptx"
    slides.save(file)
    output = import_document(file, tmp_path / "library")
    obj = load(output)
    assert any("Systole" in b["text"] and b["location"] == "slide 1" for b in obj["blocks"])
    assert any("Aorta | Vessel" in b["text"] for b in obj["blocks"])
    assert obj["figures"][0]["location"] == "slide 1"
    assert (output / obj["figures"][0]["path"]).exists()


def test_word_paragraph_table_and_images(tmp_path):
    image = tmp_path / "anatomy.png"
    make_image(image)
    doc = Document()
    doc.add_heading("Anatomy", 1)
    doc.add_paragraph("Anatomical position matters")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Superior"
    table.cell(0, 1).text = "Above"
    doc.add_picture(str(image), width=Inches(1))
    file = tmp_path / "notes.docx"
    doc.save(file)
    output = import_document(file, tmp_path / "library")
    obj = load(output)
    assert any("Anatomical position" in b["text"] for b in obj["blocks"])
    assert any("Superior | Above" in b["text"] for b in obj["blocks"])
    assert len(obj["figures"]) == 1
    assert any("positions" in w for w in obj["warnings"])


def test_cli_import_inspect_and_invalid_file(tmp_path, capsys):
    notes = tmp_path / "cells.txt"
    notes.write_text("Red blood cells carry oxygen\n")
    assert main(["import", str(notes), "--library", str(tmp_path / "library")]) == 0
    folder = next((tmp_path / "library").iterdir())
    assert main(["inspect", str(folder)]) == 0
    assert "Text blocks: 1" in capsys.readouterr().out
    assert main(["import", str(tmp_path / "nonexistent.pdf")]) == 1


def test_reject_unsupported_extension(tmp_path):
    file = tmp_path / "lecture.exe"
    file.write_bytes(b"no")
    with pytest.raises(ValueError, match="Unsupported"):
        import_document(file, tmp_path / "library")
