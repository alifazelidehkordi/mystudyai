"""Import course documents with traceable locations and embedded figure assets.

This module performs extraction only. It does not infer medical facts, run OCR,
search the internet, or transmit course materials to a model or service.
"""

from __future__ import annotations

import hashlib
import json
import mimetypes
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".docx", ".md", ".txt"}
SAFE_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "tiff", "tif", "svg"}


@dataclass
class NoteBlock:
    location: str
    text: str


@dataclass
class Figure:
    location: str
    path: str  # relative to the imported document's directory
    sha256: str
    mime_type: str


@dataclass
class ImportedDocument:
    source_name: str
    source_sha256: str
    source_type: str
    blocks: list[NoteBlock] = field(default_factory=list)
    figures: list[Figure] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_stem(name: str) -> str:
    stem = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return stem[:55] or "lecture"


def _add_figure(raw: bytes, ext: str, location: str, output: Path, figures: list[Figure]) -> None:
    if not raw:
        return
    ext = ext.lower().lstrip(".")
    if ext == "jpeg":
        ext = "jpg"
    if ext not in SAFE_IMAGE_EXTENSIONS:
        ext = "bin"
    checksum = hashlib.sha256(raw).hexdigest()
    image_name = f"{checksum[:20]}.{ext}"
    image_dir = output / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    image_path = image_dir / image_name
    if not image_path.exists():
        image_path.write_bytes(raw)
    media_type = mimetypes.guess_type(image_name)[0] or "application/octet-stream"
    figures.append(Figure(location, f"images/{image_name}", checksum, media_type))


def _pdf(source: Path, output: Path, doc: ImportedDocument) -> None:
    import fitz  # PyMuPDF

    with fitz.open(source) as pdf:
        if pdf.needs_pass:
            raise ValueError("Password-protected PDF; unlock it before importing")
        for page_index, page in enumerate(pdf, start=1):
            location = f"page {page_index}"
            text = page.get_text("text", sort=True).strip()
            if text:
                doc.blocks.append(NoteBlock(location, text))
            else:
                doc.warnings.append(f"{location}: no selectable text; scanned pages need OCR")
            for image in page.get_text("dict").get("blocks", []):
                if image.get("type") == 1:
                    _add_figure(image.get("image", b""), image.get("ext", "bin"), location, output, doc.figures)


def _pptx(source: Path, output: Path, doc: ImportedDocument) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(source)

    def collect_shapes(shapes: Iterable, location: str) -> None:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                collect_shapes(shape.shapes, location)
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    doc.blocks.append(NoteBlock(location, text))
            if shape.has_table:
                for row in shape.table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        doc.blocks.append(NoteBlock(location, " | ".join(values)))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                _add_figure(shape.image.blob, shape.image.ext, location, output, doc.figures)

    for index, slide in enumerate(presentation.slides, start=1):
        collect_shapes(slide.shapes, f"slide {index}")
        if not any(block.location == f"slide {index}" for block in doc.blocks):
            doc.warnings.append(f"slide {index}: no selectable text; images may need manual review")


def _docx(source: Path, output: Path, doc: ImportedDocument) -> None:
    from docx import Document
    from docx.table import Table

    word = Document(source)
    for index, item in enumerate(word.iter_inner_content(), start=1):
        location = f"element {index}"
        if isinstance(item, Table):
            for row in item.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    doc.blocks.append(NoteBlock(location, " | ".join(values)))
        else:
            text = item.text.strip()
            if text:
                doc.blocks.append(NoteBlock(location, text))
    # Word doesn't store reliable physical page numbers in a .docx file.
    for relation in word.part.rels.values():
        if relation.reltype.endswith("/image") and not relation.is_external:
            part = relation.target_part
            ext = Path(str(part.partname)).suffix
            _add_figure(part.blob, ext, "document (position not tracked)", output, doc.figures)
    if doc.figures:
        doc.warnings.append("DOCX figure positions are not yet tied to specific paragraphs")


def _text(source: Path, doc: ImportedDocument) -> None:
    for index, line in enumerate(source.read_text(encoding="utf-8-sig").splitlines(), start=1):
        if line.strip():
            doc.blocks.append(NoteBlock(f"line {index}", line.rstrip()))


def render_markdown(doc: ImportedDocument) -> str:
    """Create a reviewable note with explicit citations back to the source file."""
    lines = [f"# {doc.source_name}", "", f"Source SHA-256: `{doc.source_sha256}`", ""]
    locations = list(dict.fromkeys([x.location for x in doc.blocks] + [x.location for x in doc.figures]))
    for location in locations:
        lines.extend([f"## {location.capitalize()}", ""])
        for block in doc.blocks:
            if block.location == location:
                lines.extend([block.text, "", f"*Source: {doc.source_name}, {location}*", ""])
        for figure in doc.figures:
            if figure.location == location:
                if figure.mime_type.startswith("image/"):
                    lines.append(f"![Embedded figure from {location}]({figure.path})")
                else:
                    lines.append(f"[Embedded figure file]({figure.path})")
                lines.extend([f"*Source: {doc.source_name}, {location}*", ""])
    if doc.warnings:
        lines.extend(["## Import warnings", ""])
        lines.extend(f"- {warning}" for warning in doc.warnings)
        lines.append("")
    return "\n".join(lines)


def import_document(source: str | Path, library: str | Path = "library") -> Path:
    """Import a file and return a folder containing document.json, notes.md and images/.

    Content-hash directories avoid filename collisions and make re-imports safe.
    """
    source = Path(source).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Not a file: {source}")
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported file type '{suffix}'; choose {', '.join(sorted(SUPPORTED_SUFFIXES))}")

    checksum = _sha256_file(source)
    output = Path(library).expanduser().resolve() / f"{_safe_stem(source.stem)}-{checksum[:12]}"
    output.mkdir(parents=True, exist_ok=True)
    doc = ImportedDocument(source.name, checksum, suffix.lstrip("."))
    try:
        if suffix == ".pdf":
            _pdf(source, output, doc)
        elif suffix == ".pptx":
            _pptx(source, output, doc)
        elif suffix == ".docx":
            _docx(source, output, doc)
        else:
            _text(source, doc)
    except Exception:
        # Avoid leaving partially imported content that looks complete.
        import shutil
        shutil.rmtree(output)
        raise
    if not doc.blocks:
        doc.warnings.append("No extractable text found. Image-only material requires OCR or manual transcription.")
    (output / "document.json").write_text(json.dumps(asdict(doc), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output / "notes.md").write_text(render_markdown(doc), encoding="utf-8")
    return output
